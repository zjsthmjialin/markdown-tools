"""
Extract embedded images from PDF, DOCX, and PPTX files.
"""

import os


class ImageExtractor:
    """Static utility class for extracting embedded images from various document formats."""

    @staticmethod
    def extract_from_pdf(file_path, output_dir):
        """Extract images from a PDF file using PyMuPDF (fitz).

        Args:
            file_path: Path to the PDF file.
            output_dir: Directory where an ``images/`` subfolder will be created.

        Returns:
            dict: Mapping of {internal_name: saved_filename}. Empty dict on error or
                when no images are found.
        """
        try:
            import fitz
        except ImportError:
            return {}

        images_dir = os.path.join(output_dir, 'images')
        try:
            os.makedirs(images_dir, exist_ok=True)
        except OSError:
            return {}

        result = {}
        try:
            doc = fitz.open(file_path)
            try:
                seen_xrefs = set()
                for page_idx in range(len(doc)):
                    page = doc[page_idx]
                    img_list = page.get_images(full=True)
                    for img_idx, img_info in enumerate(img_list):
                        xref = img_info[0]
                        if xref in seen_xrefs:
                            continue
                        seen_xrefs.add(xref)

                        try:
                            extracted = doc.extract_image(xref)
                        except Exception:
                            continue

                        if not extracted:
                            continue

                        ext = extracted.get('ext', 'png')
                        image_bytes = extracted.get('image')
                        if not image_bytes:
                            continue

                        filename = f'pdf_page{page_idx + 1}_img{img_idx + 1}.{ext}'
                        filepath = os.path.join(images_dir, filename)

                        # Avoid overwriting: append suffix if needed
                        filepath, filename = ImageExtractor._unique_path(filepath, images_dir)

                        with open(filepath, 'wb') as f:
                            f.write(image_bytes)

                        internal_name = f'xref_{xref}'
                        result[internal_name] = filename
            finally:
                doc.close()
        except Exception:
            pass

        return result

    @staticmethod
    def extract_from_docx(file_path, output_dir):
        """Extract images from a DOCX file using python-docx.

        Args:
            file_path: Path to the DOCX file.
            output_dir: Directory where an ``images/`` subfolder will be created.

        Returns:
            dict: Mapping of {original_ref: saved_filename}. Empty dict on error or
                when no images are found.
        """
        try:
            from docx import Document
        except ImportError:
            return {}

        images_dir = os.path.join(output_dir, 'images')
        try:
            os.makedirs(images_dir, exist_ok=True)
        except OSError:
            return {}

        result = {}
        try:
            document = Document(file_path)
            img_counter = 0
            for rel in document.part.rels.values():
                if 'image' not in rel.reltype:
                    continue

                img_counter += 1
                try:
                    blob = rel.target_part.blob
                except Exception:
                    continue

                if not blob:
                    continue

                ext = ImageExtractor._ext_from_ref(rel.target_ref)
                filename = f'docx_image_{img_counter}.{ext}'
                filepath = os.path.join(images_dir, filename)
                filepath, filename = ImageExtractor._unique_path(filepath, images_dir)

                with open(filepath, 'wb') as f:
                    f.write(blob)

                result[rel.target_ref] = filename
        except Exception:
            pass

        return result

    @staticmethod
    def extract_from_pptx(file_path, output_dir):
        """Extract images from a PPTX file using python-pptx.

        Args:
            file_path: Path to the PPTX file.
            output_dir: Directory where an ``images/`` subfolder will be created.

        Returns:
            dict: Mapping of {original_ref: saved_filename}. Empty dict on error or
                when no images are found.
        """
        try:
            from pptx import Presentation
        except ImportError:
            return {}

        images_dir = os.path.join(output_dir, 'images')
        try:
            os.makedirs(images_dir, exist_ok=True)
        except OSError:
            return {}

        result = {}
        try:
            prs = Presentation(file_path)
            img_counter = 0
            seen_targets = set()
            for slide in prs.slides:
                for rel in slide.part.rels.values():
                    if 'image' not in rel.reltype:
                        continue

                    # Deduplicate across slides (same image may appear multiple times)
                    target_ref = getattr(rel, 'target_ref', None) or getattr(rel, 'target_partname', None)
                    if target_ref and str(target_ref) in seen_targets:
                        continue
                    if target_ref:
                        seen_targets.add(str(target_ref))

                    img_counter += 1
                    try:
                        blob = rel.target_part.blob
                    except Exception:
                        continue

                    if not blob:
                        continue

                    ext = ImageExtractor._ext_from_ref(str(target_ref) if target_ref else '')
                    filename = f'pptx_image_{img_counter}.{ext}'
                    filepath = os.path.join(images_dir, filename)
                    filepath, filename = ImageExtractor._unique_path(filepath, images_dir)

                    with open(filepath, 'wb') as f:
                        f.write(blob)

                    result[str(target_ref) if target_ref else f'image_{img_counter}'] = filename
        except Exception:
            pass

        return result

    # ------------------------------------------------------------------
    # Internal helpers
    # ------------------------------------------------------------------

    @staticmethod
    def _unique_path(filepath, images_dir):
        """Return (filepath, filename) that does not collide with existing files."""
        if not os.path.exists(filepath):
            return filepath, os.path.basename(filepath)

        base, ext = os.path.splitext(os.path.basename(filepath))
        counter = 1
        while True:
            new_name = f'{base}_{counter}{ext}'
            new_path = os.path.join(images_dir, new_name)
            if not os.path.exists(new_path):
                return new_path, new_name
            counter += 1

    @staticmethod
    def _ext_from_ref(ref, default='png'):
        """Guess a file extension from a relationship target reference string."""
        if not ref:
            return default
        # Take the last path segment
        name = ref.rsplit('/', 1)[-1]
        # Common image extensions found in OOXML packages
        for known in ('png', 'jpg', 'jpeg', 'gif', 'bmp', 'tiff', 'tif', 'emf', 'wmf', 'svg'):
            if known in name.lower():
                return known if known != 'jpeg' else 'jpg'
        # Fallback: try to extract whatever comes after the last dot
        if '.' in name:
            ext = name.rsplit('.', 1)[-1].lower()
            if ext.isalpha() and len(ext) <= 5:
                return ext
        return default
